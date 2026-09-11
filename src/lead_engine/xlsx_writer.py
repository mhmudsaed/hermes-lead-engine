"""Enriched .xlsx output via openpyxl (plus run report sheet)."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from lead_engine.csv_input import RowError
from lead_engine.enrich import CompanyResult, RunReport

COLUMNS = [
    ("company", "Company"),
    ("domain", "Domain"),
    ("status", "Status"),
    ("homepage_url", "Homepage URL"),
    ("careers_url", "Careers URL"),
    ("about_url", "About URL"),
    ("contact", "Contact (published only)"),
    ("tech_signals", "Tech signals"),
    ("summary", "Summary"),
    ("outreach_draft", "Outreach draft"),
    ("screenshot_path", "Screenshot"),
    ("source_urls", "Source URLs"),
    ("reason", "Reason / notes"),
]

COLUMN_WIDTHS = {
    "company": 22, "domain": 24, "status": 12, "homepage_url": 34,
    "careers_url": 34, "about_url": 34, "contact": 30, "tech_signals": 36,
    "summary": 50, "outreach_draft": 60, "screenshot_path": 30,
    "source_urls": 44, "reason": 40,
}


def result_row(result: CompanyResult) -> list[str]:
    """Flatten one result to a spreadsheet row (column order = COLUMNS)."""
    return [
        result.company,
        result.domain,
        result.status,
        result.homepage_url,
        result.careers_url,
        result.about_url,
        result.contact,
        "; ".join(result.tech_signals),
        result.summary,
        result.outreach_draft,
        result.screenshot_path,
        "; ".join(result.source_urls),
        result.reason or result.notes,
    ]


def write_workbook(
    path: str,
    results: list[CompanyResult],
    report: RunReport,
    skipped: list[RowError] | None = None,
) -> str:
    """Write the enriched workbook. Returns *path*."""
    workbook = Workbook()
    leads = workbook.active
    assert leads is not None  # a fresh Workbook always has an active sheet
    leads.title = "leads"
    header_font = Font(bold=True)
    for col_idx, (_key, title) in enumerate(COLUMNS, start=1):
        cell = leads.cell(row=1, column=col_idx, value=title)
        cell.font = header_font
        leads.column_dimensions[get_column_letter(col_idx)].width = COLUMN_WIDTHS.get(
            _key, 24
        )
    for row_idx, result in enumerate(results, start=2):
        for col_idx, value in enumerate(result_row(result), start=1):
            leads.cell(row=row_idx, column=col_idx, value=value)
    leads.freeze_panes = "A2"
    leads.auto_filter.ref = leads.dimensions

    summary = workbook.create_sheet("run_report")
    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 90
    summary.append(["Metric", "Value"])
    for cell in summary[1]:
        cell.font = header_font
    summary.append(["processed", report.processed])
    summary.append(["enriched", report.enriched])
    summary.append(["failed", report.failed])
    summary.append(["blocked", report.blocked])
    summary.append(["skipped (invalid input rows)", report.skipped])
    summary.append(["reasons", ""])
    for reason in report.reasons:
        summary.append(["", reason])
    if skipped:
        summary.append(["", ""])
        summary.append(["skipped input rows", ""])
        for err in skipped:
            summary.append([f"line {err.line_no}", err.reason])
    workbook.save(path)
    return path
