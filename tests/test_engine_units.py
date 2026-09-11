"""Unit tests: xlsx writer, email/summary helpers, fail-one-continue-others."""
from __future__ import annotations

from openpyxl import load_workbook

from lead_engine.browser import FakeAdapter, PageSnapshot
from lead_engine.csv_input import InputRow
from lead_engine.enrich import (
    CompanyResult,
    EngineConfig,
    enrich_companies,
    enrich_one,
    extract_emails,
    resolve_start_url,
    summarize,
)
from lead_engine.xlsx_writer import write_workbook
from lead_engine.enrich import RunReport


def _snap(url: str, text: str = "hello world", links=None) -> PageSnapshot:
    return PageSnapshot(url=url, ok=True, status=200, title="T", text=text,
                        links=links or [])


def test_extract_emails_filters_junk():
    text = ("Write to jobs@acme.com or noreply@acme.com, "
            "also test@example.com and Jobs@Acme.com twice.")
    assert extract_emails(text) == ["jobs@acme.com"]


def test_summarize_picks_first_substantial_sentence():
    text = "Hi. Acme Robotics builds autonomous warehouse robots for busy ports."
    summary = summarize(text)
    assert "warehouse robots" in summary
    assert summarize("") == ""


def test_resolve_start_url_modes():
    assert resolve_start_url("https://acme.com/x") == "https://acme.com/x"
    assert resolve_start_url("acme.com") == "https://acme.com"
    assert resolve_start_url("127.0.0.1:8607/acme") == "http://127.0.0.1:8607/acme"
    assert (resolve_start_url("acme-robotics", "http://h:8607/{slug}/")
            == "http://h:8607/acme-robotics/")


def test_fail_one_site_continues_others(tmp_path):
    base = "http://fake"
    good = _snap(
        f"{base}/good/", "Good Corp makes widgets with WordPress.",
        links=[(f"{base}/good/about", "About"), (f"{base}/good/contact", "Contact")],
    )
    pages = {
        f"{base}/good/": good,
        f"{base}/good/about": _snap(f"{base}/good/about", "About Good Corp, founded 2020."),
        f"{base}/good/contact": _snap(
            f"{base}/good/contact", "Email sales@good.example for quotes."),
        f"{base}/bad/": PageSnapshot(url=f"{base}/bad/", ok=False,
                                     error="navigation failed: boom"),
    }
    browser = FakeAdapter(pages)
    config = EngineConfig(screenshots_dir=str(tmp_path / "shots"))
    rows = [InputRow(2, "Good Corp", "good"), InputRow(3, "Bad Corp", "bad")]
    # FakeAdapter ignores domains; patch start URLs via template below.
    config.base_url_template = base + "/{slug}/"
    results, report = enrich_companies(
        [InputRow(r.line_no, r.company, {"Good Corp": "good", "Bad Corp": "bad"}[r.company])
         for r in rows],
        browser, config,
    )
    assert report.processed == 2
    assert report.enriched == 1 and report.failed == 1
    by_name = {r.company: r for r in results}
    assert by_name["Good Corp"].status == "enriched"
    assert by_name["Bad Corp"].status == "failed"
    assert "unreachable" in by_name["Bad Corp"].reason
    # isolation: separate site tags visited
    assert {(u, t) for u, t in browser.visited} >= {
        (f"{base}/good/", "good-corp"), (f"{base}/bad/", "bad-corp")}
    assert browser.closed is False  # engine does not own the adapter lifecycle


def test_empty_result_stays_failed_not_guessed(tmp_path):
    pages = {"http://fake/empty/": _snap("http://fake/empty/", "Hi.")}
    browser = FakeAdapter(pages)
    config = EngineConfig(base_url_template="http://fake/{slug}/",
                          screenshots_dir=str(tmp_path / "s"))
    result = enrich_one(InputRow(2, "Empty", "empty"), browser, config)
    # "Hi." is too thin for a summary, no signals/contact/careers -> failed,
    # and nothing may be guessed into the empty fields.
    assert result.status == "failed"
    assert result.careers_url == "" and result.contact == ""
    assert result.tech_signals == [] and result.summary == ""


def test_write_workbook_roundtrip(tmp_path):
    out = str(tmp_path / "enriched.xlsx")
    results = [
        CompanyResult(company="Acme", domain="acme.com", status="enriched",
                      careers_url="https://acme.com/careers",
                      tech_signals=["React [framework signal]"],
                      screenshot_path="shots/acme.png"),
        CompanyResult(company="Bad", domain="bad", status="failed",
                      reason="homepage unreachable: boom"),
    ]
    report = RunReport(processed=2, enriched=1, failed=1,
                       reasons=["Bad: failed — boom"])
    write_workbook(out, results, report, skipped=[])
    book = load_workbook(out)
    assert book.sheetnames == ["leads", "run_report"]
    leads = book["leads"]
    assert leads.cell(1, 1).value == "Company"
    assert leads.cell(2, 1).value == "Acme"
    assert leads.cell(2, 5).value == "https://acme.com/careers"
    assert leads.cell(3, 3).value == "failed"
    summary = book["run_report"]
    values = {summary.cell(r, 1).value: summary.cell(r, 2).value
              for r in range(1, summary.max_row + 1)}
    assert values["processed"] == 2 and values["enriched"] == 1
